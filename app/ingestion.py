"""Bounded, side-effect free listing imports and public, DNS-pinned feeds."""
from __future__ import annotations
import csv
from .construction import construction_status
from .sale_scope import rental_listing
import hashlib
import http.client
import io
import ipaddress
import json
import math
import socket
import ssl
import time
import zipfile
from datetime import datetime
from urllib.parse import urlsplit
import xml.etree.ElementTree as ET

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
ALIASES = {
    'source': ('origem',), 'external_id': ('_id', 'id'), 'url': ('link_id',),
    'price': ('salePrice',), 'address': ('streetFullName',),
    'neighborhood': ('regionName', 'neighbourhood'), 'parking': ('parkingSpaces', 'parkingSpots'),
    'bathrooms': ('restrooms',), 'amenities': ('installations',),
    'combined_monthly_cost': ('iptuPlusCondominium',),
    'condo_fee': ('complexFee',),
    'metro_minutes': ('walkingTimeToMetro',), 'image_url': ('coverImage',),
    'metro_station': ('nearestMetroStation',),
    'property_type': ('type',),
}
NUMBERS = ('price', 'area', 'bedrooms', 'bathrooms', 'parking', 'floor', 'floor_min_reported', 'floor_max_reported', 'condo_fee',
           'property_tax', 'combined_monthly_cost', 'metro_minutes', 'latitude', 'longitude')
TEXT = ('title', 'address', 'neighborhood', 'city')

def _missing(value):
    return value is None or (isinstance(value, str) and value.strip().lower() in ('', 'null', 'none', 'nan', 'unknown', 'n/a'))

def _number(value, field):
    if _missing(value):
        return None
    if isinstance(value, bool):
        raise ValueError(f'{field}: número inválido')
    if isinstance(value, str):
        value = value.strip().replace('R$', '').replace(' ', '')
        if ',' in value:
            value = value.replace('.', '').replace(',', '.')
    try:
        number = float(value)
    except (ValueError, TypeError):
        raise ValueError(f'{field}: número inválido') from None
    if not math.isfinite(number):
        raise ValueError(f'{field}: número não finito')
    if field not in ('floor', 'floor_min_reported', 'floor_max_reported', 'latitude', 'longitude') and number < 0:
        raise ValueError(f'{field}: valor negativo')
    if field in ('bedrooms', 'bathrooms', 'parking', 'floor', 'floor_min_reported', 'floor_max_reported') and number != int(number):
        raise ValueError(f'{field}: deve ser inteiro')
    if field == 'latitude' and not -90 <= number <= 90 or field == 'longitude' and not -180 <= number <= 180:
        raise ValueError(f'{field}: coordenada inválida')
    return number

def _url(value):
    if _missing(value):
        return None
    value = str(value).strip()
    if len(value) > 2048:
        raise ValueError('URL excede 2048 caracteres')
    parsed = urlsplit(value)
    if parsed.scheme not in ('http', 'https') or not parsed.hostname or parsed.username or parsed.password or any(ord(c) < 33 for c in value):
        raise ValueError('URL inválida: use HTTP(S) sem credenciais')
    try:
        parsed.port
    except ValueError:
        raise ValueError('Porta inválida') from None
    return value

def _normalize(raw):
    if not isinstance(raw, dict):
        raise ValueError('Anúncio deve ser um objeto')
    data = {str(k).strip(): v for k, v in raw.items() if k is not None}
    explicit_combined_cost = not _missing(data.get('combined_monthly_cost'))
    for field, aliases in ALIASES.items():
        if _missing(data.get(field)):
            for alias in aliases:
                if not _missing(data.get(alias)):
                    data[field] = data[alias]
                    break
    if rental_listing(data):
        raise ValueError('Anúncios de aluguel não são aceitos: este sistema é exclusivo para compra.')
    record = {key: _number(data.get(key), key) for key in NUMBERS}
    record['transaction_type'] = 'sale'
    lower, upper = record['floor_min_reported'], record['floor_max_reported']
    if lower is not None and upper is not None and lower > upper:
        raise ValueError('Intervalo de andar invertido')
    record['combined_cost_period'] = (
        data.get('combined_cost_period') if data.get('combined_cost_period') in ('monthly', 'unknown')
        else 'monthly'
    ) if explicit_combined_cost else 'unknown'
    for field in ('price', 'area'):
        if record[field] is None or record[field] <= 0:
            raise ValueError(f'{field}: obrigatório e maior que zero')
    record.update({key: None if _missing(data.get(key)) else str(data[key]).strip()[:2000] for key in TEXT})
    record['url'] = _url(data.get('url'))
    record['image_url'] = _url(data.get('image_url'))
    source = data.get('source')
    if _missing(source):
        source = urlsplit(record['url']).hostname if record['url'] else 'import'
    record['source'] = str(source).strip()
    if len(record['source']) > 200:
        raise ValueError('source excede 200 caracteres')
    external = data.get('external_id')
    if _missing(external):
        if not record['url']:
            raise ValueError('Informe external_id ou URL para identificar o anúncio')
        external = hashlib.sha256(record['url'].encode()).hexdigest()[:32]
    record['external_id'] = str(int(external)) if isinstance(external, float) and external.is_integer() else str(external).strip()
    if len(record['external_id']) > 500:
        raise ValueError('external_id excede 500 caracteres')
    observed = data.get('observed_at')
    if isinstance(observed, datetime):
        observed = observed.isoformat()
    if not _missing(observed):
        try:
            observed = datetime.fromisoformat(str(observed).replace('Z', '+00:00')).isoformat()
        except ValueError:
            raise ValueError('observed_at: data ISO inválida') from None
    else:
        observed = None
    record['observed_at'] = observed
    record['construction_status'] = construction_status(data)
    record['tax_period'] = data.get('tax_period') if data.get('tax_period') in ('monthly', 'annual') else 'unknown'
    kind = str(data.get('property_type') or '').lower()
    record['property_type'] = {'apartamento': 'apartment', 'apartment': 'apartment', 'residential / apartment': 'apartment', 'casa': 'house', 'house': 'house', 'studio': 'studio', 'kitnet': 'studio'}.get(kind)
    if not record['title']:
        label = {'apartment': 'Apartamento', 'house': 'Casa', 'studio': 'Studio'}.get(record['property_type'], 'Imóvel')
        area_label = format(record['area'], 'g').replace('.', ',')
        record['title'] = f'{label} de {area_label} m²'
        if record['neighborhood']:
            record['title'] += ' · ' + record['neighborhood']
    for field, choices in {'condition': ('good', 'needs_work'), 'sunlight': ('good', 'poor'), 'ventilation': ('good', 'poor'), 'documentation': ('verified', 'pending')}.items():
        record[field] = data.get(field) if data.get(field) in choices else 'unknown'
    sale = data.get('forSale', data.get('status'))
    record['status'] = 'active' if str(sale).lower() in ('true', '1', 'active', 'available', 'published', 'for_sale') else 'unavailable' if str(sale).lower() in ('false', '0', 'unavailable', 'sold', 'inactive') else 'unknown'
    elevator = data.get('elevator')
    record['elevator'] = True if str(elevator).lower() in ('true', '1', 'sim') else False if str(elevator).lower() in ('false', '0', 'não', 'nao') else None
    occupied = str(data.get('occupied')).strip().lower()
    record['occupied'] = True if occupied in ('true', '1', 'sim') else False if occupied in ('false', '0', 'não', 'nao') else None
    station = data.get('metro_station')
    if not _missing(station) and (not isinstance(station, str) or len(station.strip()) > 150):
        raise ValueError('metro_station: informe texto de até 150 caracteres')
    record['metro_station'] = None if _missing(station) else station.strip()
    amenities = data.get('amenities')
    if isinstance(amenities, str):
        try:
            amenities = json.loads(amenities)
        except ValueError:
            amenities = [v.strip() for v in amenities.split(';') if v.strip()]
    record['amenities'] = [str(v)[:200] for v in amenities][:100] if isinstance(amenities, list) else []
    record['data_origin'] = 'import'
    record['provenance'] = {k: {'source': record['source'], 'observed_at': observed} for k, v in record.items() if v is not None}
    return record

def _xml_rows(content):
    if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper() or b'\x00' in content:
        raise ValueError('XML DTD/entidades ou codificação não suportada')
    root = ET.fromstring(content)
    mapping = {'ListingID': 'external_id', 'Title': 'title', 'ListPrice': 'price', 'LivingArea': 'area', 'Bedrooms': 'bedrooms', 'Bathrooms': 'bathrooms', 'Garage': 'parking', 'PropertyAdministrationFee': 'condo_fee', 'YearlyTax': 'property_tax', 'City': 'city', 'Neighborhood': 'neighborhood', 'Address': 'address', 'PropertyType': 'property_type', 'DetailViewUrl': 'url', 'LastUpdateDate': 'observed_at', 'Latitude': 'latitude', 'Longitude': 'longitude'}
    for listing in root.iter():
        if listing.tag.split('}')[-1] != 'Listing':
            continue
        row = {'source': 'vrsync', 'tax_period': 'annual'}
        for element in listing.iter():
            name = element.tag.split('}')[-1]
            if name in mapping and element.text:
                row[mapping[name]] = element.text.strip()
            if name == 'Item' and element.attrib.get('medium') == 'image' and 'image_url' not in row:
                row['image_url'] = element.text
        yield row

def parse_upload(content: bytes, filename: str) -> dict:
    result = {'records': [], 'errors': [], 'warnings': []}
    try:
        if not isinstance(content, bytes) or len(content) > MAX_BYTES:
            raise ValueError('Arquivo excede limite de 10 MB')
        extension = filename.lower().rsplit('.', 1)[-1]
        if extension == 'xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024 or any(e.file_size > 20 * 1024 * 1024 or e.file_size / max(e.compress_size, 1) > 1000 for e in entries):
                    raise ValueError('XLSX excede limite de descompressão')
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1 or sheet.max_column and sheet.max_column > 200:
                    raise ValueError('Planilha excede limite de 5000 linhas/200 colunas')
                values = sheet.iter_rows(values_only=True)
                headers = next(values, ())
                rows = []
                for index, row in enumerate(values):
                    if index >= MAX_ROWS:
                        raise ValueError('Limite de 5000 linhas excedido')
                    if any(v is not None for v in row):
                        rows.append(dict(zip(headers, row)))
            finally:
                workbook.close()
        elif extension == 'json':
            rows = json.loads(content.decode('utf-8-sig'))
            if isinstance(rows, dict):
                rows = rows.get('listings')
            if not isinstance(rows, list):
                raise ValueError('JSON deve conter uma lista ou {listings: [...]}')
        elif extension == 'xml':
            rows = _xml_rows(content)
        elif extension == 'csv':
            decoded = content.decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            rows = csv.DictReader(io.StringIO(decoded), dialect=dialect)
        else:
            raise ValueError('Formato aceito: CSV, JSON, XLSX ou XML VRSync')
        for index, row in enumerate(rows, start=1):
            if index > MAX_ROWS:
                raise ValueError('Limite de 5000 anúncios excedido')
            try:
                result['records'].append(_normalize(row))
            except (ValueError, TypeError, OverflowError) as error:
                result['errors'].append({'row': index, 'message': str(error)})
        if any(r['observed_at'] is None for r in result['records']):
            result['warnings'].append('Data de observação desconhecida preservada; importação não comprova anúncio atual.')
        if any(r['combined_monthly_cost'] is not None for r in result['records']):
            result['warnings'].append('Custo combinado IPTU + condomínio preservado sem inventar sua divisão.')
        result['warnings'].append('Importação não comprova cobertura completa do catálogo da fonte.')
    except (ValueError, TypeError, UnicodeError, csv.Error, zipfile.BadZipFile, ET.ParseError, KeyError, OSError) as error:
        result['records'] = []
        result['errors'].append({'row': 0, 'message': str(error)})
    return result

def _resolve_public(url):
    parsed = urlsplit(_url(url) or '')
    if not parsed.hostname:
        raise ValueError('URL de feed obrigatória')
    port = parsed.port or (443 if parsed.scheme == 'https' else 80)
    if port not in (80, 443):
        raise ValueError('Feed deve usar porta 80 ou 443')
    addresses = socket.getaddrinfo(parsed.hostname, port, type=socket.SOCK_STREAM)
    if not addresses:
        raise ValueError('DNS sem endereço')
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if (not ip.is_global or ip.is_multicast or ip.is_reserved or ip.is_unspecified
                or getattr(ip, 'ipv4_mapped', None) and not ip.ipv4_mapped.is_global
                or getattr(ip, 'sixtofour', None) and not ip.sixtofour.is_global
                or getattr(ip, 'teredo', None)):
            raise ValueError('Feed deve apontar exclusivamente para endereços públicos')
    return parsed, addresses[0][4][0], port

def _download_once(url):
    parsed, address, port = _resolve_public(url)
    connection = http.client.HTTPConnection(address, port, timeout=12)
    try:
        connection.connect()
        if parsed.scheme == 'https':
            connection.sock = ssl.create_default_context().wrap_socket(connection.sock, server_hostname=parsed.hostname)
        path = parsed.path or '/'
        if parsed.query:
            path += '?' + parsed.query
        connection.request('GET', path, headers={'Host': parsed.netloc, 'Accept-Encoding': 'identity', 'User-Agent': 'ImovelRadar/1.0 authorized-feed'})
        response = connection.getresponse()
        if 300 <= response.status < 400:
            raise ValueError('Redirecionamento de feed bloqueado; configure a URL final')
        if response.status != 200:
            raise ValueError(f'Feed retornou HTTP {response.status}')
        length = response.getheader('Content-Length')
        if length and int(length) > MAX_BYTES:
            raise ValueError('Feed excede limite de 10 MB')
        if response.getheader('Content-Encoding', 'identity') != 'identity':
            raise ValueError('Feed comprimido não suportado')
        chunks, size = [], 0
        deadline = time.monotonic() + 30
        while True:
            if time.monotonic() > deadline:
                raise TimeoutError('Tempo total do feed excedido')
            chunk = response.read(min(65536, MAX_BYTES + 1 - size))
            if not chunk:
                break
            chunks.append(chunk)
            size += len(chunk)
            if size > MAX_BYTES:
                raise ValueError('Feed excede limite de 10 MB')
        return b''.join(chunks), response.getheader('Content-Type', '')
    finally:
        connection.close()

def fetch_feed(url: str) -> dict:
    for attempt in range(2):
        try:
            content, content_type = _download_once(url)
            break
        except (OSError, http.client.HTTPException):
            if attempt:
                raise ValueError('Falha de rede ao obter feed; tente novamente') from None
            time.sleep(0.25)
    extension = 'json' if 'json' in content_type or content.lstrip().startswith((b'[', b'{')) else 'xml' if 'xml' in content_type or content.lstrip().startswith(b'<') else 'csv'
    result = parse_upload(content, 'feed.' + extension)
    for record in result['records']:
        record['data_origin'] = 'feed'
        if record['source'] in ('import', 'vrsync'):
            record['source'] = 'feed:' + hashlib.sha256(url.encode()).hexdigest()[:24]
            for provenance in record['provenance'].values():
                provenance['source'] = record['source']
    return result
